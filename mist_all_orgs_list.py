#!/usr/bin/env python3
"""List every org on every Mist cloud, with its MSP, to a dated Excel sheet.

Uses the same mist_api_usage_tracker.ini as the usage tracker: one
'cloud = api_token' line per cloud under [clouds], plus the optional [network]
section for Zscaler / proxies. The tokens need super-user (/api/v1/super)
access on their cloud.

Step 1 builds an MSP table (id -> name, tier) for every cloud. Step 2 pages
through /super/stats/orgs on every cloud and writes one row per org to
mist_orgs_YYYY-MM-DD.xlsx next to this script.

Before collecting anything, every cloud's /super access is checked. If any
cloud fails, the failures are listed and the script stops. Any failure later
in the run also stops it, so the spreadsheet is only written when every cloud is
complete.
"""
import math
import os
import re
import sys
import time
from datetime import datetime, timezone
from urllib.parse import urlparse

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Shares the .ini, cloud parsing and proxy-aware HTTP session with the tracker.
from mist_api_usage_tracker import INI_NAME, SCRIPT_DIR, SESSION, load_clouds

PAGE_LIMIT = 1000
REQUEST_TIMEOUT = 60
MAX_RETRIES = 5
MAX_COLUMN_WIDTH = 30
DATE_FORMAT = "yyyy-mm-dd hh:mm:ss"
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")

COLUMNS = [
    "Org Name", "ID", "Cloud", "No. Sites", "Inventory", "Connected",
    "Created EPOC", "Created Time", "Modified EPOC", "Modified Time",
    "MSP_ID", "MSP_Name", "MSP_Tier", "URL", "Session Expiry", "Num Devices",
    "Num Devices Disconnected", "Allow Mist",
]


def cloud_name(host):
    """api.mist.com -> 'us', api.eu.mist.com -> 'eu', api.gc3.mist.com -> 'gc3'."""
    region = host[len("api."):-len("mist.com")].rstrip(".")
    return region or "us"


def epoch_to_datetime(epoch):
    """Epoch seconds -> naive UTC datetime, which Excel stores as a real date."""
    return datetime.fromtimestamp(epoch, tz=timezone.utc).replace(tzinfo=None) if epoch else None


def clean_string(value):
    """Strip control characters that Excel can't store."""
    if isinstance(value, str):
        return re.sub(r"[\x00-\x08\x0B-\x0C\x0E-\x1F]", "", value)
    return value


class CloudFailure(Exception):
    """A request to a cloud failed - the script stops rather than write a
    spreadsheet with that cloud missing or incomplete."""


def get(url, token):
    """GET with rate-limit retries. Raises CloudFailure if the request can't
    be made at all."""
    headers = {"Authorization": f"Token {token}"}
    for attempt in range(1, MAX_RETRIES + 1):
        try:
            resp = SESSION.get(url, headers=headers, timeout=REQUEST_TIMEOUT)
        except requests.exceptions.SSLError as e:
            sys.exit(f"\nTLS certificate verification failed for {url}: {e}\n"
                     f"If you're behind a TLS-inspecting proxy (e.g. Zscaler), set [network] "
                     f"ca_bundle in {INI_NAME} to the path of its root CA certificate (PEM format).")
        except requests.exceptions.ProxyError as e:
            sys.exit(f"\nCould not reach the proxy for {url}: {e}\n"
                     f"Check [network] http_proxy / https_proxy in {INI_NAME} "
                     f"or your HTTP(S)_PROXY environment variables.")
        except requests.exceptions.RequestException as e:
            raise CloudFailure(f"could not connect ({e.__class__.__name__})")
        if resp.status_code == 429:
            wait = int(resp.headers.get("Retry-After", 30))
            print(f"    ⏳ Rate limited - waiting {wait}s (attempt {attempt}/{MAX_RETRIES})...")
            time.sleep(wait)
            continue
        return resp
    raise CloudFailure(f"still rate limited after {MAX_RETRIES} attempts")


def explain_failure(resp):
    if resp.status_code == 401:
        return "HTTP 401 - token not valid on this cloud"
    if resp.status_code == 403:
        return "HTTP 403 - token has no super-user (/super) access"
    return f"HTTP {resp.status_code}: {resp.text[:200]}"


def token_owner(host, token):
    """Email of the user a token belongs to, for failure messages."""
    try:
        resp = get(f"https://{host}/api/v1/self", token)
    except CloudFailure:
        return None
    return resp.json().get("email") if resp.ok else None


def check_super_access(clouds):
    """Make one small /super request per cloud. Returns [(cloud, reason)] for
    every cloud that failed."""
    failures = []
    for host, token in clouds:
        try:
            resp = get(f"https://{host}/api/v1/super/stats/orgs?limit=1", token)
            if resp.ok:
                print(f"  ✅ {cloud_name(host)}")
                continue
            reason = explain_failure(resp)
            if resp.status_code == 403:
                owner = token_owner(host, token)
                if owner:
                    reason += f" (token belongs to {owner})"
        except CloudFailure as e:
            reason = str(e)
        print(f"  ❌ {cloud_name(host)}")
        failures.append((cloud_name(host), reason))
    return failures


def fetch_paged(base_url, token):
    """Yield (page_items, page, total_pages) for a page-based Mist list
    endpoint. X-Page-Total is the total item count, not the page count.
    Raises CloudFailure on any error response."""
    page = 1
    while True:
        sep = "&" if "?" in base_url else "?"
        resp = get(f"{base_url}{sep}limit={PAGE_LIMIT}&page={page}", token)
        if not resp.ok:
            raise CloudFailure(f"{explain_failure(resp)} on {urlparse(base_url).path} page {page}")
        items = resp.json()
        if not items:
            return
        total_items = int(resp.headers.get("X-Page-Total", len(items)))
        total_pages = max(1, math.ceil(total_items / PAGE_LIMIT))
        yield items, page, total_pages
        if page >= total_pages or len(items) < PAGE_LIMIT:
            return
        page += 1


def fetch_msps(host, token):
    """Return {msp_id: {"name": ..., "tier": ...}} for one cloud."""
    msps = {}
    for items, _, _ in fetch_paged(f"https://{host}/api/v1/super/msps", token):
        for m in items:
            if m.get("id"):
                msps[m["id"]] = {"name": m.get("name"), "tier": m.get("tier")}
    return msps


def all_org_stats(host, token):
    """Yield the stats of every org on one cloud."""
    for orgs, page, total_pages in fetch_paged(f"https://{host}/api/v1/super/stats/orgs", token):
        print(f"  📄 Page {page}/{total_pages} - collected {len(orgs)} orgs")
        yield from orgs


def org_rows(host, token, msps):
    """Yield one row dict per org on one cloud."""
    name = cloud_name(host)
    manage = "manage.mist.com" if name == "us" else f"manage.{name}.mist.com"
    for org in all_org_stats(host, token):
        org_id = org.get("id")
        msp_id = org.get("msp_id")
        msp = msps.get(msp_id, {})
        yield {
            "Org Name": org.get("name"),
            "ID": org_id,
            "Cloud": name,
            "No. Sites": org.get("num_sites"),
            "Inventory": org.get("num_inventory"),
            "Connected": org.get("num_devices_connected"),
            "Created EPOC": org.get("created_time"),
            "Created Time": epoch_to_datetime(org.get("created_time")),
            "Modified EPOC": org.get("modified_time"),
            "Modified Time": epoch_to_datetime(org.get("modified_time")),
            "MSP_ID": msp_id,
            "MSP_Name": msp.get("name"),
            "MSP_Tier": msp.get("tier"),
            "URL": f"https://{manage}/admin/?org_id={org_id}",
            "Session Expiry": org.get("session_expiry"),
            "Num Devices": org.get("num_devices"),
            "Num Devices Disconnected": org.get("num_devices_disconnected"),
            "Allow Mist": org.get("allow_mist"),
        }


def print_summary(summary):
    print(f"\n{'Cloud':<8}{'MSPs':>8}{'Orgs':>10}{'Sites':>10}{'Devices':>10}{'Connected':>11}")
    print("-" * 57)
    totals = [0, 0, 0, 0, 0]
    for name, counts in summary.items():
        print(f"{name:<8}" + "".join(f"{v:>{w},}" for v, w in zip(counts, (8, 10, 10, 10, 11))))
        totals = [t + v for t, v in zip(totals, counts)]
    print("-" * 57)
    print(f"{'Total':<8}" + "".join(f"{v:>{w},}" for v, w in zip(totals, (8, 10, 10, 10, 11))))


def cell_text_length(value):
    if value is None:
        return 0
    if isinstance(value, datetime):
        return len(DATE_FORMAT)
    return len(str(value))


def write_workbook(filename, rows):
    """One 'Orgs' sheet: styled, frozen header row, auto filter on every
    column, and each column sized to its longest value (capped at 30)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Orgs"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")

    widths = [len(c) for c in COLUMNS]
    date_cols = [i + 1 for i, c in enumerate(COLUMNS) if c in ("Created Time", "Modified Time")]
    for row in rows:
        values = [clean_string(row[c]) for c in COLUMNS]
        ws.append(values)
        widths = [max(w, cell_text_length(v)) for w, v in zip(widths, values)]
    for col in date_cols:
        for (cell,) in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            cell.number_format = DATE_FORMAT

    for i, width in enumerate(widths, start=1):
        # +2 leaves room for the filter drop-down arrow.
        ws.column_dimensions[get_column_letter(i)].width = min(width + 2, MAX_COLUMN_WIDTH)
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    wb.save(filename)


def print_failures(failures):
    print(f"\n🛑 Stopped - {len(failures)} cloud(s) failed. No spreadsheet was written.\n")
    print(f"{'Cloud':<8}Reason")
    print("-" * 70)
    for name, reason in failures:
        print(f"{name:<8}{reason}")
    sys.exit(1)


def collect(clouds):
    """Steps 1 and 2. Returns (rows, summary); raises CloudFailure (with the
    cloud name set on it) as soon as any cloud fails."""
    # --- Step 1: MSP table for every cloud ---
    print("\n🔍 Building MSP table across all clouds...")
    all_msps = {}
    for host, token in clouds:
        print(f"  Fetching MSPs from {cloud_name(host)}...")
        try:
            all_msps[host] = fetch_msps(host, token)
        except CloudFailure as e:
            e.cloud = cloud_name(host)
            raise
        print(f"    ✅ Found {len(all_msps[host])} MSPs in {cloud_name(host)}")

    # --- Step 2: orgs for every cloud ---
    print("\n🚀 Starting org collection across all clouds...")
    rows, summary = [], {}
    for host, token in clouds:
        name = cloud_name(host)
        print(f"\n📋 Processing orgs from {name}...")
        try:
            cloud_rows = list(org_rows(host, token, all_msps[host]))
        except CloudFailure as e:
            e.cloud = name
            raise
        rows.extend(cloud_rows)
        summary[name] = [
            len(all_msps[host]),
            len(cloud_rows),
            sum(r["No. Sites"] or 0 for r in cloud_rows),
            sum(r["Num Devices"] or 0 for r in cloud_rows),
            sum(r["Connected"] or 0 for r in cloud_rows),
        ]
    return rows, summary


def main():
    clouds = load_clouds()

    # --- Step 0: every cloud must have super-user access ---
    print("🔐 Checking super-user access on every cloud...")
    failures = check_super_access(clouds)
    if failures:
        print_failures(failures)

    try:
        rows, summary = collect(clouds)
    except CloudFailure as e:
        print_failures([(e.cloud, str(e))])

    # --- Step 3: save ---
    filename = os.path.join(SCRIPT_DIR, f"mist_orgs_{datetime.now():%Y-%m-%d}.xlsx")
    print(f"\n💾 Saving {len(rows)} total orgs to file...")
    write_workbook(filename, rows)

    print_summary(summary)
    print(f"\n✅ Excel exported: {filename}")
    print("🎉 Done!")


if __name__ == "__main__":
    main()
